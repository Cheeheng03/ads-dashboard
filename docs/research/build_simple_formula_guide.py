from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont

OUT=Path(__file__).resolve().parent
QA=OUT/'formula_guide_qa'; QA.mkdir(exist_ok=True)
d=Document(); s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=Inches(.62); s.bottom_margin=Inches(.62)
s.left_margin=Inches(.75); s.right_margin=Inches(.75)
s.header_distance=Inches(.25); s.footer_distance=Inches(.27)
for name,size in [('Normal',11.5),('Title',29),('Subtitle',13),('Heading 1',22),('Heading 2',14)]:
 st=d.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(8); st.paragraph_format.line_spacing=1.08
 if name.startswith('Heading'): st.paragraph_format.space_before=Pt(12)
for st in d.styles:
 for node in list(st.element.iter(qn('w:pBdr'))): node.getparent().remove(node)
d.core_properties.title='Ad Observatory Guide to Spend and Performance'
d.core_properties.subject='Plain language formulas and examples for the local dashboard'
d.core_properties.author='Prepared for Tang Chee Heng'
hp=s.header.paragraphs[0]; hp.text='AD OBSERVATORY     /     PLAIN LANGUAGE GUIDE'
for r in hp.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor(0,0,0)
fp=s.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=fp.add_run('Local dashboard • 4 October 2026     |     '); r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)

def p(t,bold=False):
 pp=d.add_paragraph(); rr=pp.add_run(t); rr.bold=bold; return pp

def h(t): d.add_heading(t,2)
def page(t): d.add_page_break(); d.add_heading(t,1)
def eq(t):
 pp=d.add_paragraph(); pp.paragraph_format.space_before=Pt(3); pp.paragraph_format.space_after=Pt(9)
 math=OxmlElement('m:oMath'); run=OxmlElement('m:r'); prop=OxmlElement('m:rPr'); sty=OxmlElement('m:sty'); sty.set(qn('m:val'),'p'); prop.append(sty); prop.append(OxmlElement('m:nor')); run.append(prop)
 tx=OxmlElement('m:t'); tx.text=t; run.append(tx); math.append(run); pp._p.append(math)
 return pp

def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers)); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for col,w in zip(t.columns,widths): col.width=Inches(w)
 pr=t._tbl.tblPr; borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'5'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
 pr.append(borders)
 for cells,vals,header,idx in [(t.rows[0].cells,headers,True,0)]+[(t.add_row().cells,row,False,i) for i,row in enumerate(rows)]:
  trpr=cells[0]._tc.getparent().get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
  if header: trpr.append(OxmlElement('w:tblHeader'))
  for j,(cell,txt,w) in enumerate(zip(cells,vals,widths)):
   cell.width=Inches(w); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcp=cell._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
   for side,val in [('top',90),('bottom',90),('left',110),('right',110)]:
    n=OxmlElement('w:'+side); n.set(qn('w:w'),str(val)); n.set(qn('w:type'),'dxa'); mar.append(n)
   tcp.append(mar); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EFF2' if header else ('F6F8F9' if idx%2 else 'FFFFFF')); tcp.append(sh)
   pp=cell.paragraphs[0]; pp.paragraph_format.space_after=Pt(0); pp.paragraph_format.line_spacing=1.04
   if j and len(str(txt))<23: pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
   r=pp.add_run(str(txt)); r.font.size=Pt(10.5); r.bold=header; r.font.color.rgb=RGBColor(0,0,0)
 d.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

def link(label,url):
 pp=d.add_paragraph(); pp.paragraph_format.space_after=Pt(6)
 hl=OxmlElement('w:hyperlink'); hl.set(qn('r:id'),pp.part.relate_to(url,RT.HYPERLINK,is_external=True))
 rr=OxmlElement('w:r'); props=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'145C85'); props.append(color); sz=OxmlElement('w:sz'); sz.set(qn('w:val'),'21'); props.append(sz); rr.append(props)
 tx=OxmlElement('w:t'); tx.text=label; rr.append(tx); hl.append(rr); pp._p.append(hl)

# Original educational chart. No external image assets.
fontpath='/System/Library/Fonts/Supplemental/Arial.ttf'
def chart():
 im=Image.new('RGB',(1500,440),'white'); dr=ImageDraw.Draw(im)
 f=ImageFont.truetype(fontpath,29); small=ImageFont.truetype(fontpath,25)
 dr.text((30,12),'A fixed RM100 daily pool shared by two ads',font=f,fill='#15242B')
 for label,y,a,b,caption in [('Days 1 to 15',105,100,0,'A gets RM100 per day'),('Days 16 to 30',250,50,50,'A and B each get RM50 per day')]:
  dr.text((30,y+22),label,font=f,fill='#15242B')
  x=340; total=1050
  dr.rectangle((x,y,x+total*a/100,y+63),fill='#276D66')
  if b: dr.rectangle((x+total*a/100,y,x+total,y+63),fill='#6B78AF')
  dr.text((x+15,y+14),'Ad A',font=f,fill='white')
  if b: dr.text((x+total*a/100+15,y+14),'Ad B',font=f,fill='white')
  dr.text((x,y+77),caption,font=small,fill='#38484E')
 dr.text((30,390),'Illustration only. Meta does not promise equal allocation.',font=small,fill='#38484E')
 path=QA/'shared_budget.png'; im.save(path); return path

# PAGE 1
D='Ad Observatory Guide to Spend and Performance'
d.add_heading(D,0)
p('Simple explanations and worked examples for your local dashboard')
p('Prepared for Chee Heng • 4 October 2026')
h('What this guide helps you understand')
p('You want to know how much a business spends, which ads have been around longest, and which ads get a good response. This guide explains what we can calculate, where the inputs come from, and what the answers actually mean.')
p('Think of an ad as a poster for a cookie shop. We can see the poster and its recorded start date. Seeing the poster does not let us see the shop’s advertising bill or how many people bought cookies.')
p('Our calculator can answer “What would the spending be if these assumptions were true?” Real spending and performance need supporting records. A correct formula cannot make an unknown input true.',True)
h('The labels you should look for')
table(['Label','Plain meaning'],[
 ('Observed','Saved from a source at a particular check. The source may be incomplete or old.'),
 ('Supplied reported','You entered a number from a report. The app does not independently verify it.'),
 ('Calculated','Arithmetic using inputs. Its reliability depends on those inputs.'),
 ('Assumed or estimated','A what if result using one or more guesses or benchmarks.'),
 ('Unknown','The input is missing or incompatible. This does not mean zero.')],[1.55,5.45])
h('Where to start')
p('Pages 2–3 explain Meta and the vocabulary. Pages 4–8 explain the formulas. Pages 9–10 explain comparisons and charts. Page 11 shows how to enter and save data. Pages 12–13 give a formula reference and sources.')
p('Every ringgit example in this guide is fictional. The default ranges are teaching inputs, not measured Malaysian prices or facts about Amazin’ Graze or JOBBIE.')

# PAGE 2
page('1 How Meta charges for ads')
p('Imagine giving Meta a spending instruction for your cookie shop. You choose a goal, such as website visits or purchases, and a budget. Meta’s delivery system decides which eligible ads to show within the campaign settings.')
h('Budget and actual spend are different')
p('A budget says how much you intend to allow. Spend says how much delivery actually used. A daily budget is an average spending instruction; actual daily delivery can vary. Budgets can sit at campaign or ad set level, with several ads sharing them. [1]')
p('An ad set is a group of ads with shared delivery settings. Our dashboard simplifies this into one assumed company pool. That is our modeling choice, not a budget discovered inside the competitor’s account.')
h('Why there is no price per poster per day')
p('In auction buying, opportunities to show an ad are competed for. Meta describes advertiser bids, the estimated chance of the desired action, and ad quality as factors in the auction. [2]')
p('There is no universal daily price tier that tells us “one image costs RM10” or “one video costs RM30.” Format and placement can affect results and costs. They do not reveal a competitor’s bill.')
p('The goal you optimize for and the event used for billing are separate settings. Meta’s SDK lists both. Paying for delivery while optimizing for purchases does not mean every account pays a fixed fee per purchase. Supported settings vary. [3]')
h('What our spending estimate includes')
p('We model media spend: money for ad delivery. Video production, designer fees, agency fees and taxes are separate costs unless your supplied report explicitly includes them. There is no separate “maintenance fee” we can infer from an ad remaining visible.')
h('The most important example')
eq('RM100 per day × 30 days = RM3,000')
p('If that is one shared daily pool, adding more ads does not automatically increase RM3,000. It can divide the same money between more ads. Actual allocation may be very unequal; the equal split in our tool is an assumption.')

# PAGE 3
page('2 The words behind the numbers')
p('These words describe different steps. Keeping them separate prevents a video play or a like from accidentally becoming a paid click.')
table(['Word','Think of it this way'],[
 ('Impressions','How many times the ad was displayed. One person can generate several displays.'),
 ('Reach','How many distinct people or accounts were reached, using the source’s definition and reporting window.'),
 ('Frequency','Impressions ÷ reach. 200 displays to 100 people means an average frequency of 2.'),
 ('Paid link clicks','Recorded clicks on an ad’s link. Use this exact definition, not all clicks or public likes.'),
 ('Landing page visits','Recorded visits after the destination loads. A link click does not guarantee a loaded page.'),
 ('Conversions','A defined result, such as a purchase or lead. Choose one definition and keep it consistent.'),
 ('Public engagement','The same reaction, comment and share counters on a public post. May include organic activity.'),
 ('Copy and creative','Copy is the wording. Creative includes the image, video and other visible ad content.'),
 ('Scope','The boundaries of the data: dates, geography, objective, attribution and metric definitions.')],[1.7,5.3])
h('Percent just means out of one hundred')
eq('1% = 1 out of 100 = 0.01')
p('A 1% click rate means one link click for every 100 impressions on average. It does not necessarily mean one unique person clicked out of 100 unique people.')
h('Three small math signs')
p('× means multiply. ÷ means divide or share equally. A range such as RM100–RM300 gives a low and high assumption. It is not a promise that the real number falls between them.')
p('A reporting window is the date range you selected. An attribution window is a different thing: the rule for when a result is credited to an ad after an interaction.')

# PAGE 4
page('3 Spending from an assumed budget')
p('Choose “Assumed company budget” when you want to explore a scenario without delivery reports. You supply the daily pool. The app does not discover it.')
eq('Ad share for one date = daily pool ÷ eligible ads')
eq('Ad window spend = all its daily shares added together')
p('An eligible ad is saved, was Active at the last check, and has a valid recorded start date on or before that date. The model assumes continuous delivery, full use of the daily pool, and equal allocation. It excludes ads last marked Inactive or Unknown and ads with unusable start dates.')
d.add_picture(str(chart()),width=Inches(7))
table(['Fictional 30 day scenario','Ad A','Ad B','Company'],[
 ('First 15 dates','RM1,500','RM0*','RM1,500'),
 ('Next 15 dates','RM750','RM750','RM1,500'),
 ('Whole selected window','RM2,250','RM750','RM3,000')],[3.1,1.3,1.3,1.3])
p('*Ad B has not started in this example. This is a calculation within the scenario, not an observed zero from Meta.')
h('What changes when you add more ads')
p('If ten ads were eligible for all 30 dates, the same RM100 daily pool would give each RM300. The company total would still be RM3,000. If you assign a larger pool to Company B, its larger estimate comes from that assumption.')
p('When no saved ads are eligible, or the pool is blank, the tool leaves a gap. It does not assume the real company spent nothing. Using today’s last observed status to allocate past dates is also a scenario assumption.')

# PAGE 5
page('4 Spending from delivery evidence')
p('Choose “Delivery evidence & benchmarks” when you have a report with spend, impressions, reach or paid link clicks. A benchmark is a reference price from other data; it is not necessarily this ad’s price.')
h('The tool picks the first available method')
table(['Priority','Inputs available','Formula','Fictional result'],[
 ('1','Supplied spend','Use the supplied amount','RM1,500 → RM1,500'),
 ('2','Impressions and assumed CPM','Impressions × CPM ÷ 1,000','100,000 × RM10–20 ÷ 1,000\n= RM1,000–2,000'),
 ('3','Reach, assumed frequency and CPM','Reach × frequency × CPM ÷ 1,000','10,000 × 1.5–3 × RM10–20 ÷ 1,000\n= RM150–600'),
 ('4','Paid link clicks and assumed CPC','Paid link clicks × CPC','1,000 × RM0.50–2\n= RM500–2,000')],[.65,1.85,2.25,2.25])
p('CPM means cost per 1,000 impressions. CPC means cost per paid link click in this tool. In method 2, 100,000 impressions are 100 bundles of 1,000. At RM10–20 per bundle, the estimate is RM1,000–2,000.')
h('Why we choose one method')
p('These are alternative ways to estimate the same bill. Adding RM1,000–2,000 and RM500–2,000 would count the same ad twice. The expanded ad card shows alternatives so you can inspect them; the total uses only the first available method.')
p('Reported spend takes priority even when the supplied value is zero. “Reported” means supplied as a report value. It is not a claim that we have verified it with Meta.')
h('Where the missing inputs come from')
p('An authorized Ads Manager export, permitted account access, or a credible advertiser disclosure can provide evidence. Meta’s SDK requires permissioned access. A competitor’s public Page ID does not grant that access. [4]')
p('The saved Malaysian competitor collection currently has no verified spend or paid click evidence. Public creative collection alone cannot fill those fields. Some transparency datasets expose ranges for particular regions or categories; record the exact coverage before using them. [5]')

# PAGE 6
page('5 Ranges and modeled clicks')
h('How the low and high ends are calculated')
p('All spend and delivery inputs here are nonnegative. To multiply ranges, multiply low by low and high by high. To divide, use the largest denominator for the low result and the smallest denominator for the high result.')
eq('Low impressions = low spend ÷ high CPM × 1,000')
eq('High impressions = high spend ÷ low CPM × 1,000')
p('Suppose scenario spend is RM1,000–2,000 and assumed CPM is RM10–20. The low result is 1,000 ÷ 20 × 1,000 = 50,000 impressions. The high result is 2,000 ÷ 10 × 1,000 = 200,000 impressions.')
h('Turning displays into a click scenario')
eq('Modeled clicks = impressions × assumed CTR ÷ 100')
p('CTR means click through rate. With 50,000–200,000 modeled impressions and an assumed CTR of 0.5–2%, the result is 250–4,000 modeled clicks. The range gets wide because we combine several uncertain inputs.')
p('These are hypothetical clicks. They cannot be used as proof of how well the ad performed. If we assumed the click rate, the resulting clicks merely repeat that assumption.',True)
h('Which inputs each spending mode uses')
table(['Mode','Where modeled impressions come from'],[
 ('Assumed company budget','Allocated spend ÷ assumed CPM × 1,000.'),
 ('Delivery evidence & benchmarks','Supplied impressions; otherwise supplied reach × assumed frequency. In this mode, spend alone does not generate impressions.')],[2.35,4.65])
p('If you have exactly 100,000 impressions and assume 0.5–2% CTR, the click scenario is 500–2,000. Actual paid clicks remain a separate number.')
h('The range depends on your assumptions')
p('It is not a 95% confidence interval, a most likely prediction, or a guaranteed boundary for actual spending. The tool makes no probability claim and does not hide uncertainty behind an exact looking midpoint.')
p('Changing a benchmark source label does not verify the numbers. Prefer evidence matching the market, dates, audience, objective and click definition, and keep the source reference.')

# PAGE 7
page('6 Measuring a real response')
p('The following fictional report uses the same ad, dates and scope: RM1,500 spend; 100,000 exact impressions; 1,000 paid link clicks; 800 landing page visits; 20 purchases.')
table(['Metric','Formula and worked example','Meaning'],[
 ('Link CTR','1,000 ÷ 100,000 × 100 = 1%','One paid link click per 100 impressions.'),
 ('Link CPC','RM1,500 ÷ 1,000 = RM1.50','Average cost for one paid link click.'),
 ('CPA','RM1,500 ÷ 20 = RM75','Average ad cost per defined conversion; here, a purchase.'),
 ('Visit CVR','20 ÷ 800 × 100 = 2.5%','Reported conversions per 100 landing page visits under the chosen definition.')],[1.15,3.15,2.7])
p('CPA means cost per acquisition or action. CVR means conversion rate. Our evidence form uses landing page visits as the CVR denominator. A different tool may use clicks or sessions, producing a different number.')
h('Why paid clicks cannot simply be scraped')
p('The public ad and its link do not expose the advertiser’s paid click report. We can save those clicks when an authorized report or disclosure supplies them. My earlier suggestion that competitor paid clicks could simply be tracked publicly was too broad.')
p('Opening a competitor’s link tells us where it leads. It does not tell us how many other people clicked it. Public video views, reactions and comments are different measures.')
h('Exact inputs matter')
p('The tool calculates measured CTR only when the impression lower and upper bounds match, giving one supplied count. An impression range does not produce a falsely exact measured CTR. “Exact” here means one report value, not perfect tracking of every human action.')
p('CPC and CPA use supplied spend, not the modeled spend range. Measured performance stays separate even when the spending view is set to a budget scenario.')
h('How to read zero and unknown')
p('Zero clicks and 1,000 impressions give 0% CTR. RM100 spend and zero clicks give an unavailable CPC because division by zero has no finite answer. A blank click field means we do not know; it must not become zero.')

# PAGE 8
page('7 Time and public engagement')
h('Ad age is a calendar calculation')
eq('Calendar age = selected end date − source start date')
p('An ad starting 1 September has a calendar age of 29 days on 30 September. We count elapsed whole UTC calendar days. UTC is the shared time reference the app uses for these dates. A same day start has age 0. Invalid, missing or future start dates produce “unknown.”')
p('This is age since the saved ad start, not necessarily the age of its underlying public post. It does not prove that money was spent every day. Pauses, low delivery and later restarts may be invisible. An inactive ad can still have a large calendar age.')
h('Why 29 days and 30 dates can both be right')
eq('Window dates = end date − start date + 1')
eq('Window average spend = window spend ÷ window dates')
p('1–30 September includes 30 dates. RM3,000 over that window averages RM100 per calendar date. An ad that started halfway through still uses all 30 window dates for this average. It is not its cost per actual paid day.')
p('“Assumed delivery dates” counts dates allocated money in the budget model. It is not a count of observed paid days. “Oldest ad” is the largest valid calendar age among the saved ads.')
h('Public engagement needs two snapshots')
eq('Net new engagement = end counter − start counter')
eq('Net engagement per day = change ÷ elapsed days')
p('Use the same post and the same reaction + comment + share counters. If the total was 100 on 1 September and 158 on 30 September, the net change is 58. Then 58 ÷ 29 = 2 net engagements per day.')
p('The denominator here is 29 elapsed days, not 30 inclusive dates. Same date snapshots cannot give a daily pace in this tool. A lower end count produces a negative change; removals or corrections can cause that.')
p('This pace helps compare public activity over time. It can include organic activity and does not measure paid clicks, ad profitability, or engagement per paid impression. A lifetime total divided by ad age can mix activity from before the ad started.')

# PAGE 9
page('8 Comparing businesses and ads fairly')
p('Use the same window, currency and definitions. If both companies use the same assumed pool for the same eligible dates, their modeled totals will be equal by design. Compare similar markets and objectives. A traffic ad and a purchase ad may be trying to achieve different things.')
h('More clicks does not always mean a better response')
table(['Fictional ad','Paid link clicks','Impressions','Link CTR'],[
 ('Ad A','100','10,000','1%'),('Ad B','60','2,000','3%')],[1.6,1.7,1.8,1.9])
p('Ad A has more clicks. Ad B has more clicks per display. CTR adjusts for exposure, which is usually more relevant to click response than dividing clicks by age. Clicks per day would still be affected by how much delivery the ad bought.')
p('A high CTR on a tiny sample can be unstable. Neither row proves which ad earns more money. Purchases, cost per purchase and tracking quality matter when the goal is sales.')
h('What the ranking options actually mean')
table(['Sort option','The honest interpretation'],[
 ('Oldest','Longest calendar age among saved ads. A research lead, not proof of success.'),
 ('Highest spend','Largest upper bound of the selected spend range. Overlapping ranges may not establish an order.'),
 ('Measured link CTR','Largest supplied paid link click rate within the selected evidence scope.'),
 ('Public engagement pace','Largest net public counter change per elapsed day, within the selected scope.')],[2.1,4.9])
h('Why there is no magic performance score')
p('An age plus likes plus estimated spend score would mix unrelated things and turn guesses into a ranking. The current report uses understandable measures instead. It does not claim a 90 out of 100 “winner” or a probability of success.')
p('Copy is only one part of an ad. The image, offer, audience, placement and destination also affect response. To isolate copy quality, run a controlled comparison on an account you can measure, changing the wording while keeping other conditions comparable. Competitor observation can suggest ideas to test; it cannot prove what caused their results.')

# PAGE 10
page('9 Reading totals and charts')
h('Company totals only cover contributing saved ads')
eq('Tracked spend = selected ad estimates added together')
p('If 3 of 10 saved ads contribute spend, the total covers those 3 ads. It is not the entire company’s expenditure. The company average divides that total by all selected calendar dates. Unknown ads and dates are not confirmed zero spending.')
p('In evidence mode, different scope labels cause the company sum to be withheld. Matching labels are a check, not proof that the data really matches. Dates and currency must also match. No automatic currency conversion is performed.')
table(['Chart or counter','What to read'],[
 ('Cumulative spending','Budget mode only. Add daily assumed allocations from the selected start. It is not lifetime spend unless the window covers the full modeled life.'),
 ('Weekly or monthly totals','Budget mode only. Sum included dates by Monday based week or calendar month. Edge periods can contain only part of a week or month.'),
 ('Spend range by ad','Evidence mode. Up to 8 ads, sorted by upper bound. The marker shows the low estimate and the bar end the high estimate.'),
 ('Measured CTR bars','Up to 8 ads with measured CTR in the chosen scope. These bars require real supplied clicks and exact impression counts.'),
 ('Evidence coverage','Counts matching evidence entries or ads contributing spend, according to the label. It is not a percentage of truth or of all Meta ads.')],[2.0,5.0])
h('Read the axes and the gaps')
p('A shaded spending band shows the low and high scenario, not statistical confidence. Panels for two businesses may use different scales; compare the numeric labels and units, not just bar height.')
p('A gap means the tool could not allocate a value. Later cumulative values resume the sum of allocated dates only. Weekly totals can also be partial when some dates have no allocation.')
p('One total for an evidence window cannot tell us spending on each day. That is why the evidence view shows ad ranges rather than inventing a weekly spending history.')

# PAGE 11
page('10 Entering and saving your inputs')
p('Start from Companies and open a business. Use Compare for two businesses. Both views use the same report logic; the comparison shares one date window.')
h('Try a spending scenario')
p('1. Select From and To, then choose “Apply dates.” Use 1–366 calendar dates ending no later than today.')
p('2. Choose “Assumed company budget.” Open “Assumptions,” enter low and high values, and record the source or reason. “Apply assumptions” recalculates the draft.')
p('3. Inspect the spend range and each ad’s allocation. Choose “Save report” to keep the changes. Applying an edit is not the same as saving it.')
h('Add actual evidence when you have it')
p('4. On an ad card, choose “Add / edit evidence.” Supply the source URL or report reference, dates, currency and scope. For example: MY • traffic • paid link clicks • stated attribution rule.')
p('5. Enter only supported numbers. Leave unknowns blank. For exactly 100,000 impressions, enter 100,000 in both bounds. Public engagement pace needs both snapshots from the same post.')
p('6. Choose “Apply evidence,” then “Save report.” Switch the spend basis to “Delivery evidence & benchmarks” to use the evidence method priority. Save again after changing the basis.')
h('Why a number might disappear')
p('Evidence must match the report’s full From and To dates and currency. Changing the window can exclude a saved record; the tool does not divide a lifetime total into new dates for you. Open “All formulas, performance & inputs” to see the source or exclusion reason.')
p('Each ad has one current evidence entry in this model. Applying a replacement and saving overwrites that entry. Preserve your original reports separately. Each company has one saved current model, not a library of saved scenarios.')
h('What account imports do')
p('“Authorized account reports” holds imported daily account records. These remain separate from public creative cards until there is a verified ID mapping. Identical titles do not prove two records are the same ad. Importing is not the same as connecting a live API.')
p('The current local implementation does not automatically retrieve competitors’ private metrics. It saves your assumptions and evidence in your own database. Read the saved check date before treating an ad’s status as current.')

# PAGE 12
page('11 Formula reference and starting inputs')
table(['Measure','Formula or rule'],[
 ('Spend from impressions','Impressions × assumed CPM ÷ 1,000'),
 ('Spend from reach','Reach × assumed frequency × assumed CPM ÷ 1,000'),
 ('Spend from paid clicks','Supplied paid link clicks × assumed CPC'),
 ('Budget allocation','For each eligible date: shared daily pool ÷ eligible ad count; then sum'),
 ('Modeled impressions','Budget mode: spend ÷ assumed CPM × 1,000'),
 ('Modeled clicks','Impressions × assumed link CTR ÷ 100'),
 ('Measured link CTR','Supplied paid link clicks ÷ exact impressions × 100'),
 ('Measured CPC and CPA','Supplied spend ÷ paid link clicks; spend ÷ conversions'),
 ('Visit conversion rate','Conversions ÷ landing page visits × 100'),
 ('Calendar age','Elapsed UTC dates from source start to selected end'),
 ('Window daily average','Selected window spend ÷ inclusive calendar dates'),
 ('Public engagement pace','(End counter − start counter) ÷ elapsed days')],[2.0,5.0])
h('Current illustrative defaults')
table(['Input','Low','High'],[
 ('Shared daily pool','RM100','RM300'),('CPM per 1,000 impressions','RM10','RM20'),('CPC per paid link click','RM0.50','RM2'),('Frequency','1.5','3'),('Assumed link CTR','0.5%','2%')],[4.0,1.5,1.5])
p('These defaults are not market research findings. Replace them with your own justified assumptions. Supply both bounds, or leave both blank. Low must not exceed high; CPM must be positive; frequency must be at least 1; assumed CTR must be 0–100%.')
p('A missing input or zero denominator leaves a ratio unavailable. Stored arithmetic uses full precision; displayed money may be rounded. Scenario bounds are not guarantees of the real value.')

# PAGE 13
page('12 Other terms and source notes')
h('Revenue and ROAS from the earlier screens')
p('ROAS means return on ad spend. Revenue is sales value. Attributed purchase value is the value a reporting system credits to ads under its attribution rules. It may differ from total company revenue, net sales or extra sales caused by advertising.')
eq('ROAS = attributed purchase value ÷ ad spend')
p('RM4,500 attributed value ÷ RM1,500 spend = 3× ROAS, or revenue equal to 300% of spend. It is not 300% profit. Product costs, fulfilment, returns and other expenses still matter. The focused report does not project competitor revenue or use ROAS to rank ads without evidence.')
h('Useful checks when aggregating account reports')
p('Add matching ad counts and spend once, without duplicate or overlapping records. A company CTR uses total clicks ÷ total impressions, not a simple average of ad CTRs. Company CPC uses total spend ÷ total clicks; CPA uses total spend ÷ total defined conversions.')
p('Do not add reach across ads to claim unique company reach: the same person may see several ads. Incomplete records can also make a total partial. Keep currency, attribution and click definitions consistent.')
h('What this guide is based on')
p('Dashboard behavior was checked against the local source on 4 October 2026: app/estimation-model.ts, app/focused-report.tsx, app/business-model.ts and app/spend-model.ts. These define our formulas and UI behavior; they do not verify competitor inputs. Earlier research files may describe features outside this focused report.')
link('[1] Meta budget explanation, 26 July 2023, republished by Publicnow','https://www.publicnow.com/view/0807F171279348987DA5CDD63B6E59EAFFDEF343')
link('[2] Meta Creative Strategy Study Guide, October 2023, hosted by Integris Design','https://integrisdesign.com/wp-content/uploads/2024/12/Meta-Creative-Strategy-Study-Guide.pdf')
link('[3] Meta Business SDK ad set definitions for budgets and billing events','https://github.com/facebook/facebook-python-business-sdk/blob/main/facebook_business/adobjects/adset.py')
link('[4] Meta Business SDK permissions and account access documentation','https://github.com/facebook/facebook-python-business-sdk')
link('[5] Meta Ad Library API coverage and field availability','https://br-fr.facebook.com/ads/library/api/')
p('Source check: 4 October 2026. The older Meta guides explain concepts; this document does not assert a current daily overspend cap or a universal billing option. All worked examples and ranges are teaching calculations, not source reported competitor results.')

# Validate arithmetic used by examples and save.
assert 15*100+15*50 == 2250
assert 100000*10/1000 == 1000 and 100000*20/1000 == 2000
assert 10000*1.5*10/1000 == 150 and 10000*3*20/1000 == 600
assert 1000/20*1000*.5/100 == 250
assert 2000/10*1000*2/100 == 4000
assert 1000/100000*100 == 1 and 1500/1000 == 1.5
assert 1500/20 == 75 and 20/800*100 == 2.5
assert (158-100)/29 == 2
path=OUT/'Ad_Observatory_Simple_Formula_Guide.docx'; d.save(path)
print(path)
