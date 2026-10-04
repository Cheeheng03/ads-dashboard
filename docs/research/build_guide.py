from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pathlib import Path

OUT=Path('/Users/tangcheeheng/Documents/MetaAdsResearch')
d=Document()
s=d.sections[0]
s.top_margin=s.bottom_margin=Inches(.65)
s.left_margin=s.right_margin=Inches(.75)
s.page_width=Inches(8.27); s.page_height=Inches(11.69)
normal=d.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.line_spacing=1.08
for st in d.styles:
 for border in list(st.element.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
for name,size in [('Title',28),('Heading 1',20),('Heading 2',13)]:
 st=d.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string('000000')
 st.paragraph_format.space_after=Pt(10)
d.core_properties.title='Meta Ads Competitor Analysis Pipeline'
d.core_properties.subject='Research guide and proposed product specification'
d.core_properties.author='Research prepared for Tang Chee Heng'
footer=s.footer.paragraphs[0]; footer.alignment=2
r=footer.add_run('Meta Ads Research  |  '); r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def p(t,style=None): return d.add_paragraph(t,style)
def h(t): d.add_heading(t,2)
def page(t): d.add_page_break(); d.add_heading(t,1)
def bullet(t): p(t,'List Bullet')
def table(headers,rows,widths=None):
 t=d.add_table(rows=1,cols=len(headers)); t.style='Light Shading Accent 1'
 for c,x in zip(t.rows[0].cells,headers): c.text=x
 for row in rows:
  for c,x in zip(t.add_row().cells,row): c.text=x
 for row in t.rows:
  pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(5)
    for rr in pp.runs: rr.font.size=Pt(10)
 trpr=t.rows[0]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:tblHeader'))
 p('')
 return t
def link(label,url):
 pp=d.add_paragraph(); hyp=OxmlElement('w:hyperlink'); hyp.set(qn('r:id'),pp.part.relate_to(url,RT.HYPERLINK,is_external=True))
 rr=OxmlElement('w:r'); props=OxmlElement('w:rPr'); col=OxmlElement('w:color'); col.set(qn('w:val'),'145C85'); props.append(col); rr.append(props)
 tx=OxmlElement('w:t'); tx.text=label; rr.append(tx); hyp.append(rr); pp._p.append(hyp)

d.add_heading('Meta Ads Competitor Analysis Pipeline',0)
p('Research guide and proposed product specification')
p('Prepared 3 October 2026')
h('What is worth building')
p('Build a curated intelligence workspace for the businesses you care about. Save their ads, preserve a history of changes, group related creative ideas, and identify messages worth testing in your own campaigns. The strongest value is turning scattered ad observations into a searchable, evidence-backed history.')
p('Keep competitor observations separate from connected account performance. Public ad activity can support research priorities, but it does not identify ordinary competitors’ exact budgets, sales, click-through rates or return on ad spend. A persuasive-looking dashboard must not turn missing data into invented financial results.')
h('The key correction to the original idea')
p('An ad visible for 90 days is evidence of persistence, not proof of profitability or 90 days of continuous spending. There is no universal daily ad price. The same visible lifespan can be consistent with a tiny retargeting budget, substantial prospecting spend, intermittent delivery, or a campaign that simply has not been reviewed.')
h('Recommended starting scope')
bullet('Track 10 to 20 businesses in one market and one industry. Resolve each business to its actual advertiser Pages before collecting ads.')
bullet('Use a data provider or supported official API, then store dated observations in your own database. Check collection quality before building sophisticated scoring.')
bullet('Launch pinboards, creative search, new-ad alerts, offer comparisons and longevity evidence first. Add money estimates only when their inputs can be inspected.')
bullet('Connect your own Meta account and commerce or CRM records later to measure real results and test ideas inspired by competitors.')
p('This guide explains the available evidence, advertising economics, collection choices, existing tools, a proposed pipeline and a validation plan. Public facts have numbered source references; formulas, architecture, scoring and planning budgets are proposed designs or explicitly illustrative examples.')

page('1 What you can actually measure')
table(['Data you want','Ordinary competitor ads','Connected account you control'],[
('Copy and creative','Usually observable, with variants and coverage gaps','Available within account permissions'),
('Start date and active status','Library or provider observation','Account status plus delivery reporting'),
('Exact spend and daily budget','Generally unavailable','Spend reporting and permitted budget fields'),
('Impressions and reach','Region and ad-category dependent','Reported metrics; some are estimated'),
('Clicks and CTR','Generally unavailable','Available; specify which click metric'),
('Video plays and retention','Public post views are not ad reporting','Video metrics where supported'),
('Purchases and ROAS','Unavailable without advertiser disclosure','Attributed conversions if tracking exists'),
('Net company revenue','Separate financial or modeled source','Commerce or CRM reconciliation required')])
p('Meta’s Library API currently documents political or issue ads over seven years and all ad types delivered in the UK or EU over one year. Political ads can expose spend and impression bands. UK/EU delivery adds regional transparency; EU payer information is separately scoped. This does not provide universal commercial spend disclosure. [1]')
p('Marketing API Insights is a different permissioned system. Meta’s own examples and SDK expose spend, clicks, impressions, actions, action values and ROAS-related fields. Access is tied to an authorized ad account, not merely knowing a competitor’s Page ID. [2, 3]')
h('Product labels that prevent confusion')
p('Use four labels: Observed, Reported by connected account, Modeled scenario, and Unknown. Store the source, time, market, definition and reporting window behind every value. Missing is not zero. A provider claiming “performance” may mean a longevity score rather than clicks or sales.')
p('For a Malaysia-only commercial advertiser, do not plan on the official Library API providing the complete feed. A brand headquartered in Malaysia can still have UK/EU-delivered ads; coverage follows delivery geography. Validate the requested country and exact advertiser Pages in the pilot.')

page('2 How Meta advertising costs work')
p('Meta uses an auction. Its published explanation describes advertiser bids, predicted likelihood of the desired action and ad quality as inputs to delivery. Prices therefore depend on the audience and auction opportunity, rather than a fixed price for keeping an ad visible for a day. [4]')
h('Budget is different from cost')
p('A budget is the advertiser’s spending instruction. Actual spend is what delivery consumes. Multiple ads can share a campaign or ad-set budget and receive very unequal allocations. You cannot divide a company’s visible ad count into equal daily spending units. A lifetime budget is also not a guarantee of an equal amount every day.')
table(['Term','Meaning and formula'],[
('CPM','Spend ÷ impressions × 1,000. Cost for one thousand displays.'),
('CPC','Spend ÷ clicks. Specify all clicks, link clicks or outbound clicks.'),
('CTR','Clicks ÷ impressions × 100%. Use matching click definitions.'),
('CPA','Spend ÷ attributed acquisitions. Define purchase or qualified lead.'),
('Reach and frequency','Reach counts distinct accounts or people per definition. Frequency = impressions ÷ reach for the same scope.'),
('ROAS','Attributed purchase value ÷ ad spend. Revenue multiple, not profit.')])
h('Formats are not fixed price tiers')
p('Image, carousel, video, Stories and Reels do not give you a standard price-per-day lookup table. Format, placement and creative quality can change results; producing a video may also cost more than producing an image. Production cost is separate from Meta media spend. Meta’s Reels guidance encourages testing placement-appropriate creative. [5]')
h('A benchmark is context rather than a competitor measurement')
p('WordStream’s 2026 report covers almost 1,800 campaigns and distinguishes traffic from leads objectives. Its reported average CPC is US$0.60 for traffic and US$1.80 for leads. These sample averages are not Malaysian rate cards, ecommerce purchase benchmarks or evidence of a specific advertiser’s costs. [6]')
p('For planning your own campaign, an assumed RM40 acquisition cost and a goal of five acquisitions per day implies RM200 per day of media spend. This is arithmetic using an assumption, not a recommended minimum or a forecast. Replace assumptions with your own measured results.')

page('3 Estimating cumulative ad spend honestly')
h('Use the strongest available input')
table(['Evidence','Permitted calculation','How to display it'],[
('Authorized account reporting','Sum daily spend once per ad and day','Reported spend with currency and dates'),
('Public spend bands','Preserve lower and upper bounds','Disclosed interval, not exact midpoint'),
('Actual impression measure','Impressions × assumed CPM ÷ 1,000','Modeled range for matching region and period'),
('Reach only','Reach × assumed frequency × assumed CPM ÷ 1,000','Scenario with two uncertain inputs'),
('Only ad age and status','No defensible observed money calculation','Spend unknown; optional user scenario')])
h('Worked example with hypothetical inputs')
p('Suppose a regional transparency record shows 100,000 accounts reached during a defined period. Assume frequency between 1.5 and 4 and CPM between RM10 and RM40. These numbers are invented to teach the model, not measured Malaysian benchmarks.')
p('Low scenario = 100,000 × 1.5 × RM10 ÷ 1,000 = RM1,500.\nHigh scenario = 100,000 × 4 × RM40 ÷ 1,000 = RM16,000.')
p('That more than tenfold span shows how uncertain the result is. It is a sensitivity range, not a statistical confidence interval. If the reach covers the EU, the estimate covers that EU delivery only. Expressing it in RM does not make it Malaysian spend.')
h('Why days multiplied by a daily rate fails')
p('Thirty days × an assumed RM50 equals RM1,500 only under that scenario. The ad could have received almost no allocation on many days. A platform minimum budget, where applicable, is not a minimum daily spend for each ad. Active or eligible status is not a receipt for delivery.')
h('Rules for accumulation')
bullet('Keep daily flows separate from lifetime snapshots. Repeated cumulative values must never be added together.')
bullet('Do not treat the increase in lifetime reach as daily impressions. It can represent newly reached accounts and misses repeat exposures.')
bullet('Deduplicate Library IDs across countries and providers. Preserve distinct paid ad instances even when their creative assets match.')
bullet('Sum money only across disjoint records, matching periods, currencies and scope. Preserve open-ended ranges; do not invent an upper limit.')
p('A company total also depends on identifying all regional Pages, sub-brands and partnership advertisers. Label a partial result as spend or activity for tracked Pages, never guaranteed total company spend.')

page('4 Relating ads to revenue')
h('For competitors')
p('Public financial filings may disclose company or segment revenue, but they rarely isolate Meta-attributed sales. Website revenue estimates are a different source again. Similarweb’s website dataset includes an estimated monthly ecommerce revenue metric. It should remain labeled as a vendor estimate, with domain coverage and period attached. [7]')
p('Website traffic growth alongside more observed ads is a correlation. Promotions, email, search, offline sales, pricing and seasonality can move at the same time. Similarweb’s URL Spend metric is based on paid search and refers to Google landing pages; it is not a Meta spend estimate. [8]')
p('Do not divide company revenue by modeled Meta spend and call the result Meta ROAS. If useful, display a clearly named scenario such as total online revenue divided by assumed Meta spend, together with both uncertainties. It does not establish attribution or incremental sales.')
h('For your own company or an authorized client')
p('Combine daily Meta Insights with orders, refunds and customer status from your commerce platform, or qualified opportunities and closed deals from a CRM. Shopify’s GraphQL Order object provides an authorized route to order and refund information. Revenue definitions still need to be chosen and reconciled. [2, 9]')
table(['Measure','What it tells you'],[
('Meta reported ROAS','Meta-attributed purchase value / Meta spend, under the recorded attribution settings.'),
('Net sales','A chosen commerce definition after discounts and refunds; specify tax and shipping treatment.'),
('MER','Total revenue / total advertising spend across included channels; a blended efficiency measure.'),
('New customer CAC','Acquisition spend / newly acquired customers, with a stated allocation rule.'),
('Incremental return','Additional revenue caused by advertising / additional spend; requires a causal design such as a suitable holdout.')])
p('Keep the attribution window, timezone, currency, event definition and extraction time. Re-fetch recent periods for late conversions. Deduplicate purchase events, exclude test orders and reconcile refunds; do not add two platforms’ attributed revenue as if their credited sales are disjoint.')
p('Illustration: RM10,000 of attributed sales on RM3,000 spend is 3.33× ROAS. At 30% contribution margin before ads, RM3,000 remains before ad cost, leaving zero before fixed overhead. High revenue and profitability are different questions.')

page('5 Using longevity and copy analysis')
h('Treat longevity as a reason to investigate')
p('A persistent ad may reflect a useful message. It can also be cheap retargeting, low delivery, branding, a seasonal schedule, or neglected campaign management. Ending an ad can reflect inventory or timing rather than failure. These alternative explanations cannot be ruled out by a Library card.')
p('Track source start date, first observed date, last observed active date, successful observation days, gaps and explicit inactivity separately. A failed fetch must not become an “ad stopped” event. A missing ad should remain uncertain until a successful, sufficiently complete check provides evidence.')
h('A proposed study priority score')
p('Start with transparent evidence components instead of a probability of profitability: persistence within comparable cohorts, repeated reuse of the concept, and additional related variants. Where comparable regional reach is available, show it separately. More variants can indicate testing or localization rather than scaling.')
p('A first heuristic could weight persistence 50%, concept reuse 30% and variant expansion 20%, using percentile ranks within the same brand, market, format and launch period. This is a proposed research ranking only. Show data completeness beside it, decline to score sparse cohorts, and validate weights later against authorized results. Do not label 80/100 as an 80% chance of success.')
h('What the creative analysis should extract')
bullet('Exact hook and opening video moments; problem, desired outcome, offer, price, CTA, proof and objections addressed.')
bullet('Transcript with timestamps, on-screen text, visual format, product demonstration, spokesperson and landing-page promise.')
bullet('Language and local context. Preserve Malay, English and Chinese originals and store translations separately where relevant.')
bullet('Creative families: the same concept can appear with different IDs, visuals, copy or regional wording.')
h('What a useful explanation sounds like')
p('“Observed on 24 successful checks over 30 days; the same guarantee appears in four distinct creatives. The copy may reduce purchase anxiety. A competing explanation is a persistent remarketing campaign. Test a guarantee-led hook against a product-led hook in your own account, holding the offer and landing page constant.”')
p('Separate observed evidence, plausible mechanism, alternative explanation and proposed test. A model can explain a hypothesis; it cannot prove why a competitor ad earned money. For actual results, compare qualified leads, purchases or contribution alongside CTR, and avoid attributing a whole creative’s result to its copy alone.')

page('6 Existing products worth evaluating')
p('Prices below are public USD listings checked on 3 October 2026, before applicable taxes. Vendor claims describe advertised capabilities, not an independent accuracy test. Billing periods and included collection rights matter more than headline prices.')
table(['Product','Useful for this project','Public commercial context'],[
('Foreplay [10, 11]','Pinboards, competitor tracking through Spyder, and API access. Strong candidate for buying the collection layer.','Basic $59/month; Workflow $175/month shown. Non-legacy plans list 10,000 monthly API credits; 1 ad = 1 credit.'),
('Atria [12]','Brand following, searchable creative library, creative workflow, REST API and MCP.','Core $129/month billed annually, 50 followed brands and 1,200 API calls/month. Core highlights say no AI insights for followed brands; verify plan gating.'),
('Motion [13, 14]','Competitor research plus analysis of your connected account; useful for closing the research-to-test loop.','Creative Research is advertised free. Full analytics listing starts at $250/month; MCP access is documented.'),
('AdSpy [15]','Broad Facebook and Instagram ad discovery, copy and comment search. Useful for finding ideas outside your watchlist.','$149/month advertised introductory price. Confirm export/API suitability before choosing it for ingestion.'),
('MagicBrief [16]','Creative research and competitor analysis; another product to compare for workflow usability.','Advertises estimated competitor performance. Request the definition and a current quote; estimates are not private account results.'),
('Similarweb [7, 8]','Traffic and ecommerce revenue context around a company or domain.','Dataset and API access depend on plan. Revenue estimates do not identify Meta-attributed revenue.'),
('Apify [17]','Managed extraction actor with API scheduling and exports; practical technical pilot option.','Official actor lists $5.80/1,000 results on Free, down to $3.40 on Business. Plan and enrichment costs differ.'),
('SearchAPI [18]','Structured Meta Library search and detail responses as a service.','Evaluate request billing, pagination, detail calls and country coverage; it is a third-party service, not Meta’s API.')])
p('My shortlist: evaluate Foreplay first for saving and tracking workflows, and a managed extraction provider for flexible ingestion. Atria is attractive if its research and creative features replace other subscriptions. Use Motion when connecting your own results is central. Buy only after checking the same target Pages against a manual sample.')

page('7 APIs versus browser automation')
table(['Method','Best role','Main limitation'],[
('Meta Library API','Supported programmatic access for covered ads','Country/category scope, access requirements and field availability'),
('Meta Marketing API','Performance of authorized ad accounts','Cannot retrieve private competitor account metrics'),
('Vendor API','Fast custom pipeline with structured inputs','Cost, coverage, licensing and dependency on provider'),
('Playwright','Repeatable permitted UI checks, captures and fallback workflows','Selectors, pagination, sessions and UI changes require maintenance'),
('Computer use','Exploration, verifying Page identity and inspecting unusual cards','Slow and hard to make reliable for large recurring collections')])
h('Recommended decision')
p('Use an API for the daily pipeline. Keep manual inspection or computer use for onboarding and quality checks. If a permitted browser path is needed, Playwright gives testable locators and retry behavior; use semantic locators and explicit completion checks, not fragile screen coordinates. [19]')
p('A provider’s use of the word API does not mean Meta has granted it platform access. Confirm its right to collect, store and supply the data for your intended use, including redistributing it in a future SaaS product. Meta states that automated collection without permission violates its terms. Browser automation is therefore an access decision as well as an engineering choice. [20]')
h('Official API implementation notes')
p('The Library endpoint is /ads_archive. Plan for identity/location confirmation, a developer app and token. Query exact Page IDs with country and category filters; follow pagination. Useful fields include creative text, delivery dates, snapshot URL and regional reach where supported. Validate the current schema and coverage with a live token before committing. [1]')
p('For connected accounts, use GET /act_<AD_ACCOUNT_ID>/insights with ad-level daily reporting. Candidate fields include ad_id, spend, impressions, clicks, outbound_clicks, actions and action_values. Use read-only permissions where possible. Meta distinguishes standard access for your own accounts and advanced access for other people’s accounts. [2, 21]')
h('Operational rules')
p('Pin and periodically review API versions. Store credentials server-side, redact tokens from snapshot URLs and logs, follow rate-limit signals, retry with backoff and checkpoint pagination. Record whether the result was complete, partial, blocked or empty. Do not hard-code an unverified universal requests-per-hour limit.')

page('8 Proposed pipeline and data model')
p('Watchlist → scheduled collection → raw evidence → normalized ads and observations → creative grouping → analysis → review dashboard and alerts. This architecture is a proposed design, independent of the data vendor.')
table(['Layer','What it should do'],[
('Business registry','Map a business to verified Page IDs, domains, sub-brands, regions and partnership Pages. Keep uncertain matches for review.'),
('Collectors','One adapter per provider or official API. Persist cursor, requested country, query, fetched time, counts and completion state.'),
('Evidence store','Keep raw JSON and allowed media/snapshots with hashes, source URLs and retention rules. Public URLs may expire.'),
('Normalization','Keep provider ID, Library ID and connected-account ad ID in separate namespaces. Preserve missing-field reasons.'),
('Creative analysis','OCR, transcription, language detection and visual/copy tagging. Deduplicate expensive analysis by content hash.'),
('Analytics','Compute observational timelines and comparable cohort ranks. Isolate financial scenarios from reported metrics.'),
('Delivery','Company timelines, pinboards, compare view, saved searches, evidence panels and meaningful change notifications.')])
h('Core records')
p('Business and AdvertiserPage capture identity. AdInstance captures source identity and delivery metadata. AdObservation captures a dated status under specific filters. CreativeAsset and CreativeFamily separate exact files from related concepts. CollectionRun records failures and completeness. Pin and Annotation preserve your curation.')
p('MetricObservation stores metric name, value or bounds, unit, currency, geography, period, method and source. EstimateScenario stores all assumptions and model version. ConnectedAdInsight stores account metrics in a separate permissioned area. Never assume a public Library ID equals a Marketing API ad ID.')
h('A practical stack')
p('Use a TypeScript web application with PostgreSQL, object storage and a queue-backed worker. Python is optional for media processing. A scheduled job can start collection runs; n8n can validate the workflow before a custom UI. Full-text search is sufficient initially; add vector search when semantic retrieval earns its cost.')
p('Start with one daily successful snapshot per tracked Page and market. Fetch richer media analysis only for new or changed assets. Periodically re-check known ads and history to detect reactivation. Store UTC collection times and the source reporting timezone. Keep a visible coverage panel so a broken collector cannot resemble a competitor going quiet.')

page('9 Features to prioritize')
h('First release')
bullet('Curated companies and Page mappings, tags for direct competitors versus inspiration brands, and market-specific watchlists.')
bullet('Pin ads to boards with notes, labels and source links. Search original copy, transcript, offer, product and format.')
bullet('Company timeline showing newly observed ads, continuing ads, explicit stops and uncertain missing observations.')
bullet('Side-by-side creative comparison and family grouping, with a way to correct false merges.')
bullet('A study queue ranked by persistence and repetition evidence. Display coverage and reasons for each ranking.')
bullet('Weekly change digest and alerts for new offers, reactivated concepts or unusual launch activity. Alert on meaningful changes.')
h('Next additions')
bullet('Landing-page history: price, bundle, guarantee, CTA and message consistency. Label captures by region, device and date.')
bullet('Offer and creative theme trends within your tracked pool. Call this observed share of creatives, not share of ad spend.')
bullet('Creative opportunity map: overused claims, rarely addressed objections and new combinations to test.')
bullet('Localized copy comparisons and multilingual search. Compare exact wording before translating or generalizing.')
bullet('An experiment notebook linking a competitor observation to your hypothesis, creative brief, test and actual outcome.')
bullet('Reported spend and revenue dashboards for accounts you authorize, with explicit attribution definitions.')
h('Features to defer')
p('Defer exact competitor spending claims, predicted revenue from ad age, automatic “winning ad” declarations, and any score that suggests validated profitability before it has been calibrated. Add spend scenarios only after reliable regional delivery inputs or user-provided budget assumptions exist.')
h('An example company screen')
p('Company identity and tracked Pages at the top; a dated coverage badge; creative cards with pin controls; a timeline; top recurring offers; and a comparison pane. Each ad card shows original copy, source date, observation history and research-priority reasons. A money panel says Unknown unless a disclosed figure or explicit scenario supports it.')
p('The product’s defensible advantage can be the history and judgement you accumulate: verified business identities, saved evidence, corrected creative families, local-language interpretation and experiments that connect research to your own results.')

page('10 Pilot plan and operating costs')
h('A staged validation plan')
p('Stage 1: manually identify 10 to 20 target businesses and their Pages. Compare two candidate feeds against a dated Library sample, including inactive or reactivated ads where available. This checks whether the intended market is covered before substantial implementation.')
p('Stage 2: collect for two weeks. Audit pagination, duplicate IDs, missing media, delayed updates and false stop events. Retain failed runs as failures. A one-time scrape cannot establish observation history.')
p('Stage 3: build pinboards, timelines, creative search and the evidence panel. Review at least 50 varied ads manually for OCR, transcription, tags and family matching. Keep the human corrections.')
p('Stage 4: add creative hypotheses and connected-account experiments. Evaluate whether the tool saves review time and produces better test ideas. To validate a profitability proxy, use authorized outcome data and held-out brands or time periods; survival alone is an inadequate target.')
h('Proposed acceptance criteria')
bullet('Every displayed metric has source, period, scope and evidence type; unknown spend never renders as zero.')
bullet('Successful full fetches can be distinguished from partial, blocked and legitimately empty results.')
bullet('Deduplicated coverage reaches a chosen target, for example 95% of a timestamped manual sample. This measures the sample, not guaranteed global recall.')
bullet('Missing or duplicate observations do not inflate duration, launch counts or financial totals.')
bullet('Scenario arithmetic passes edge cases: zero denominators, open-ended bands, cumulative snapshots, currencies and regional overlap.')
h('Illustrative monthly cost model')
p('If 20 companies average 50 returned ads and you refresh all of them daily for 30 days, that is 30,000 returned records. At Apify’s published Free-plan rate of $5.80 per 1,000 results, the actor component would be $174 if each returned record is billable. At $3.40 it would be $102, but that rate depends on a higher plan. Add plan commitments and enabled extras; do not double-count included usage credits. [17]')
p('With Foreplay, 30,000 billable ad credits would exceed the stated 10,000 monthly inclusion. Confirm whether repeat retrievals, detail requests and tracking calls consume credits before costing the design. [11]')
p('A provisional hosting/storage allowance of $25–100/month and AI processing allowance of $10–75/month can serve as planning placeholders for a small pilot, not vendor quotes. Video volume, storage retention and model choice can move them substantially. Instrument actual usage and set spending caps.')

page('11 Decisions before implementation')
h('The recommended path')
p('Build the curation, evidence history, comparisons and experiment workflow yourself. Evaluate buying the data feed first. This preserves flexibility while avoiding the immediate burden of keeping a browser collector alive. Retain a provider adapter boundary so changing the feed does not require rebuilding the product.')
h('Choose these inputs')
table(['Decision','Practical starting assumption'],[
('Geography and sector','One market and one business type, then expand after coverage is demonstrated.'),
('Initial business pool','10 to 20 verified businesses; include regional and partnership Pages when relevant.'),
('Primary job','Find useful creative ideas and track changes; measure your own performance separately.'),
('Data purchase','One provider trial and a manual comparison before a recurring subscription.'),
('Output cadence','Daily collection and a weekly digest; urgent alerts only for selected events.'),
('Ownership model','Personal/internal tool first. Confirm redistribution rights before selling data access.'),
('Revenue integration','Commerce orders for ecommerce, or closed deals and lead quality for service businesses.')])
h('Questions for any provider demo')
p('Can it retrieve every targeted Page in the intended countries? How fresh are statuses? What historical coverage exists before tracking starts? Can media be retained? What does its performance score measure? Which fields are observed, modeled or unavailable? Are there complete pagination markers and stable IDs?')
p('Ask for the billed unit, repeat-fetch costs, detail-call charges, rate limits, API availability on the proposed plan and export limits. Ask about internal analysis, storage duration, cancellation access and commercial redistribution separately. A user-interface subscription is not automatically a data resale licence.')
h('What counts as success')
p('After a month, you should be able to answer: What changed among my competitors? Which messages keep returning? Which offers and landing pages deserve study? What evidence supports that judgement? Which experiment should we run next, and did it improve our own business outcome?')
p('Those are achievable product promises. Exact competitor CTR, profit and total Meta spend require evidence that ordinary public creative tracking does not provide. Preserve that distinction in every screen, export and AI explanation.')

sources=[
('1','Meta Ad Library API','https://br-fr.facebook.com/ads/library/api/?source=archive-landing-page','Coverage, access steps and field availability. Recent indexed documentation includes UK and EU; older page variants differ. Validate a live response before implementation.'),
('2','Meta Marketing API ad insights example','https://www.postman.com/meta/facebook-marketing-api/request/u07tack/get-ad-insights-l1','Official Meta Postman example for performance fields.'),
('3','Meta Business SDK AdsInsights source','https://github.com/facebook/facebook-python-business-sdk/blob/main/facebook_business/adobjects/adsinsights.py','Official SDK field definitions; a field name alone does not grant access or guarantee population.'),
('4','Meta Toward fairness in personalized ads','https://about.fb.com/wp-content/uploads/2023/01/Toward_fairness_in_personalized_ads.pdf','Meta explanation of auction and pacing, including bid, predicted action rate and quality.'),
('5','Meta Reels ads guidance','https://www.facebook.com/business/ads/facebook-instagram-reels-ads','Placement and creative guidance; not a universal price schedule.'),
('6','WordStream Facebook Ads Benchmarks 2026','https://www.wordstream.com/blog/facebook-ads-benchmarks-2026','First-party publication of its campaign sample, separated by traffic and leads objectives.'),
('7','Similarweb Websites Dataset','https://developers.similarweb.com/docs/websites-dataset','Estimated monthly ecommerce revenue definition.'),
('8','Similarweb URL Spend','https://support.similarweb.com/hc/en-us/articles/6466947059089-URL-Spend','Paid-search spending methodology and Google URL scope.'),
('9','Shopify GraphQL Admin Order','https://shopify.dev/docs/api/admin-graphql/latest/objects/Order','Authorized order and refund data structure.'),
('10','Foreplay pricing','https://www.foreplay.co/pricing','Public subscription tiers and competitor tracking features.'),
('11','Foreplay API credits and pricing','https://help.foreplay.co/en/help/articles/9082255-credits-and-pricing','Credit unit and monthly inclusions.'),
('12','Atria pricing','https://www.tryatria.com/pricing','Annual billing, brand following limits, REST API and plan gating.'),
('13','Motion Creative Research','https://motionapp.com/creative-research','Advertised free research features; competitor data described as public Library data.'),
('14','Motion pricing overview and MCP','https://motionapp.com/llm-info','Public analytics pricing overview. MCP documentation linked separately below.'),
('15','AdSpy official site','https://www.adspy.com/','Search capabilities and advertised introductory monthly price.'),
('16','MagicBrief features','https://magicbrief.com/features','Competitor analysis advertised as estimated performance.'),
('17','Apify Facebook Ads Library Scraper','https://apify.com/apify/facebook-ads-scraper','Official actor listing, plan-dependent per-result pricing and enrichment extras.'),
('18','SearchAPI Meta Ad Library documentation','https://www.searchapi.org/docs/meta-ad-library-api','Third-party structured search API. Detail endpoint linked separately below.'),
('19','Playwright locators','https://playwright.dev/docs/locators','Official locator and retry behavior documentation.'),
('20','Meta How We Combat Scraping','https://about.fb.com/news/2021/04/how-we-combat-scraping/','Meta’s stated position on automated collection without permission.'),
('21','Meta Marketing API official collection','https://www.postman.com/meta/facebook-marketing-api/documentation/0zr4mes/facebook-marketing-api-mapi','Standard and advanced access distinctions.')]
page('Sources and reading')
p('Research checked 3 October 2026. Source links support public product facts. Suggested architecture, prioritization, scoring, pilot criteria and worked scenarios are recommendations, not vendor guarantees.')
for n,label,url,note in sources[:11]:
 link(n+'  '+label,url); pp=p(note); pp.paragraph_format.space_after=Pt(9); pp.runs[0].font.size=Pt(10)
page('Additional sources')
for n,label,url,note in sources[11:]:
 link(n+'  '+label,url); pp=p(note); pp.paragraph_format.space_after=Pt(9); pp.runs[0].font.size=Pt(10)
link('Motion MCP documentation','https://help.motionapp.com/en/articles/14315735-motion-mcp')
link('SearchAPI ad detail documentation','https://www.searchapi.org/docs/meta-ad-library-ad-details-api')
d.save(OUT/'Meta_Ads_Analysis_Research_Guide.docx')
print(OUT/'Meta_Ads_Analysis_Research_Guide.docx')
