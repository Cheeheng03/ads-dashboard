from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path
out=Path('/Users/tangcheeheng/Documents/MetaAdsResearch')
d=Document(); sec=d.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.7)
for n in ['Normal','Title','Heading 1','Heading 2']:
 s=d.styles[n];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
 for b in list(s.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
d.styles['Normal'].font.size=Pt(10.5)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.05
for n,size in [('Title',26),('Heading 1',19),('Heading 2',13)]:d.styles[n].font.size=Pt(size)
f=sec.footer.paragraphs[0];f.alignment=2;f.add_run('Meta Ads Metrics Catalog  |  ').font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
d.core_properties.title='Meta Ads Metrics and Retrievable Data Catalog'
d.core_properties.author='Prepared for Tang Chee Heng'
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,2)
def page(t):
 pp=d.add_heading(t,1);pp.paragraph_format.page_break_before=True
def table(rows,headers=('Data to store','Availability and interpretation')):
 t=d.add_table(rows=1,cols=2);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 t.columns[0].width=Inches(2.12);t.columns[1].width=Inches(4.65)
 pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
 pr.append(b)
 for idx,row in enumerate([headers]+rows):
  cells=t.rows[0].cells if idx==0 else t.add_row().cells
  for j,(c,txt) in enumerate(zip(cells,row)):
   c.width=Inches(2.12 if j==0 else 4.65);c.text=txt;c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),'70' if side in ['top','bottom'] else '95');x.set(qn('w:type'),'dxa');m.append(x)
   cp.append(m);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF5' if idx==0 else ('F7F9FC' if idx%2==0 else 'FFFFFF'));cp.append(sh)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.line_spacing=1
    for rr in pp.runs:rr.font.size=Pt(10);rr.bold=idx==0
  rp=cells[0]._tc.getparent().get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
  if idx==0:rp.append(OxmlElement('w:tblHeader'))
 d.add_paragraph().paragraph_format.space_after=Pt(1)
def link(label,url):
 pp=d.add_paragraph();hyp=OxmlElement('w:hyperlink');hyp.set(qn('r:id'),pp.part.relate_to(url,RT.HYPERLINK,is_external=True));rr=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'17568A');rp.append(color);rr.append(rp);tx=OxmlElement('w:t');tx.text=label;rr.append(tx);hyp.append(rr);pp._p.append(hyp)

d.add_heading('Meta Ads Metrics and Retrievable Data Catalog',0)
p('Prepared for Chee Heng on 3 October 2026')
p('This catalog is the discussion inventory for your Ad Observatory dashboard. It covers competitor ads and businesses, public online enrichment, calculations we can build, and reporting that requires an advertiser’s authorization. It describes potential fields; it does not mean that they have already been collected or implemented.')
h('What we can reliably build around')
p('Prioritize creative evidence, advertiser identity, dated observations, offer changes and searchable research notes. For ordinary Malaysia or Singapore commercial ads, exact spend, clicks, conversions and revenue are not generally public. A long-lived ad is a research signal, not a proven winner. There is no fixed daily Meta price: auction delivery determines cost. [1, 9]')
table([
('Observed','Directly visible content or a returned source field. Retain the original evidence.'),
('Reported','A platform or company disclosure, sometimes estimated or expressed as a range.'),
('Derived','Our calculation from dated observations. Coverage determines reliability.'),
('Inferred','An analyst or AI interpretation, with evidence and review status.'),
('Modeled','A scenario or provider estimate. Inputs and limitations must be visible.'),
('Authorized only','Requires access to the advertiser’s account, commerce system or CRM.'),
('Unknown','Not available, not collected, or not applicable. Never silently replace with zero.')],('Evidence label','Meaning'))
h('Geography and coverage')
p('The public Library describes active ads across Meta, political or issue archives for seven years, EU archives for one year, and UK archives for one year for ads run after 1 July 2025. Meta says first appearances and updates can take up to 24 hours. [2]')
p('Important refinement to the earlier guide: current indexed API pages differ by locale. The detailed API reference retrieved here explicitly scopes commercial API access and extra fields to EU delivery. UK public archive coverage is confirmed, but UK API field parity must be tested before promising it. A Malaysian company advertising in Europe can have EU records; this does not reveal its Malaysia-only results. [1, 2]')
p('No preferred market was supplied for this catalog, so it covers MY/SG ordinary commercial ads plus conditional international disclosures.')

page('1 Data for each individual ad')
p('Core identifiers, creative text, dates and platforms are documented by Meta. Media, CTA, destination and grouping details may need the visible ad snapshot or a collector rather than a structured official API field. Preserve arrays of variants. [1, 3]')
table([
('Ad identity','Library ID, source URL, advertiser Page ID and Page name.'),
('Text','Primary copy, headline, link caption and description; language.'),
('Lifecycle','Source start and stop dates where present; currently displayed status. Creation time is separate when provided.'),
('Distribution','Publisher platforms and countries found under the recorded search scope. A search country is not a complete targeting list.'),
('Creative format','Image, video, carousel or other displayed format; number and order of accessible cards.'),
('Media references','Image/video URLs, thumbnails, source dimensions, aspect ratio and duration when accessible. URLs can expire.'),
('Call to action','Button text and CTA type, when present.'),
('Destination','Displayed destination, resolved landing URL, domain, accessible tracking parameters and redirect chain.'),
('Variant grouping','Multiple-version label, collation/group IDs and counts where exposed. A group count is not a campaign count.'),
('Partnership disclosure','Creator/partner name, paid partnership or disclaimer label when visible. No inferred personal identity.'),
('Product and offer','Advertised product, displayed price/currency, discount, bundle, trial, shipping promise, deadline and guarantee.'),
('Public post linkage','Post URL/ID only if exposed and verifiably linked. Many ads do not provide a usable public post.'),
('Public engagement','Visible post reactions, comments, shares or view counters, if accessible. Store as post metrics, not paid ad metrics.'),
('Your curation','Pin, collection, tags, notes, owner, priority, competitor relevance and planned experiment.')])
p('The product/offer and curation rows are proposed extraction and research fields. Public engagement must retain its post scope: organic and paid exposure can mix; reposts and shared post IDs can also pool engagement across ads. These counters cannot provide CTR, conversion rate or a paid-only video view count.')
p('Library IDs and Ads Manager ad IDs are different identifier systems. Similar creative, identical copy or a shared destination does not prove two records are the same ad.')

page('2 Ad history and creative analysis')
p('These are proposed fields we calculate or annotate ourselves. They require repeated snapshots or accessible media, not additional hidden Meta performance access.')
table([
('Observation timestamps','First collected, most recent check, most recent active observation and last explicit inactive observation.'),
('Ad age','Days since the source start date. Label it elapsed age, not days of continuous delivery.'),
('Observed persistence','Active observations / successful eligible checks. Also record sample count and monitoring interval.'),
('Observed active dates','Distinct dates on which active status was seen. This is a sampled lower-bound record, not daily spend.'),
('Status transitions','Observed active/inactive transitions, missing results, reappearance and observation gaps.'),
('Creative changes','Copy, CTA, asset, offer and destination diffs between snapshots; version timestamps and evidence.'),
('Creative families','Exact text/media hashes plus reviewed similarity clusters; reuse across Library IDs and Pages.'),
('Text extraction','OCR text, video transcript, on-screen captions, language and translations. Preserve the original.'),
('Opening hook','First sentence or first seconds; question, demonstration, problem, promise or pattern interruption.'),
('Message structure','Problem, desired outcome, product benefit, proof, objections, offer and CTA sequence.'),
('Creative execution','UGC-style, testimonial, founder, demo, explainer, comparison or static offer; pacing and brand appearance time.'),
('Positioning hypotheses','Apparent audience, awareness stage, emotional appeal, differentiator and likely objection. These are interpretations.'),
('Research decision','Why investigate, alternative explanation, testable hypothesis, proposed variant and your later test outcome.')])
h('How to describe an apparent winner')
p('Use “persistent creative worth testing,” supported by the checks and variants observed. Do not label it profitable because it is old, repeated or highly liked. Retargeting, small budgets, seasonal reuse and intermittent delivery are alternative explanations.')
p('An AI review should cite the exact copy, image region or video timecode behind each claim. Store model/prompt version, extraction confidence and human corrections. A confidence score describes the analysis, not the probability the ad makes money.')
h('Useful alerts')
p('New Library ID; new creative family; changed offer; changed destination; ad reappears; repeated checks show inactivity; or collection fails. Distinguish an advertiser change from a collector outage.')

page('3 Data for each business')
p('Resolve a business to all relevant advertiser Pages, domains and regional brands before aggregating. Keep a reviewed mapping: one company can use several Pages, agencies or partners. Page transparency can expose creation, rename and merge history and some manager-country information where available. [4]')
table([
('Business identity','Brand and legal name where verified, domains, industry, products, business model, market and competitor category.'),
('Advertiser map','Page IDs, Page URLs, Instagram handles, regional Pages, parent/sub-brand links and mapping evidence.'),
('Page transparency','Creation date, prior names, merges and manager countries when disclosed. Page age is not company age.'),
('Observed inventory','Distinct Library IDs, currently reported active IDs, unique creative families and covered advertiser Pages.'),
('New ad activity','First-seen IDs and new creative families in 7/30 days. Call them discoveries unless launch dates corroborate.'),
('Stopping and reappearance','Explicitly observed inactive ads, missing ads and returning ads as separate counts.'),
('Creative mix','Share by format, platform, language, product, hook, angle, CTA and offer among collected ads.'),
('Longevity distribution','Median elapsed ad age; shares older than 30/60/90 days. Active-only samples have survivor bias.'),
('Reuse and experimentation','Ads per creative family, new copy/visual variations, family recurrence and change cadence.'),
('Offer strategy','Discount incidence, discount ranges, bundle mix, subscription/trial use and promotion calendar.'),
('Destination strategy','Unique landing pages, product pages versus lead forms or advertorials; landing-page changes.'),
('Market comparisons','Observed country/language coverage and overlaps. Visibility in a country is not spend allocation.'),
('Peer comparisons','Your company’s observed ad-count share of a fixed tracked pool. Never call this spend or impression share.'),
('Research operations','Pinned ads, open hypotheses, reviewed ads, alerts, last successful collection and field completeness.')])
h('Aggregation rules')
p('Count each Library ID once within a defined scope; retain creative-family counts separately. The platform flags on one ad can overlap. Do not add per-ad reach to claim unique business reach: the same person may see several ads. Do not sum the same cumulative spend snapshot on successive days.')
p('Store source timestamps, geography, reporting period and included Pages next to each rollup. A company with more visible ads does not necessarily spend more: budget can be concentrated in a few ads.')

page('4 Additional data available online')
p('These enrichments need separate sources and are not Meta Library metrics. Availability, pricing, access and geographic coverage must be checked for each target business.')
table([
('Website and product pages','Visible price/currency, products, variants, offers, stock labels, shipping, returns, guarantee, subscription terms and review counts. A stock label is not inventory units.'),
('Landing-page history','Headline, hero, proof, CTA, form length, pricing and checkout-entry changes from repeated page captures. No competitor conversion rate.'),
('Public social profiles','Followers, public posts, posting cadence and visible engagement. Track changes with timestamps; followers are not customers.'),
('Reviews and reputation','Public ratings, review volume, recent review cadence, themes and response time where visible. Sampling/moderation bias applies.'),
('Website audience estimates','Estimated visits, channel mix, countries, engagement duration and referral sources from providers such as Similarweb. Site-level estimates, not ad-ID outcomes. [5]'),
('Search demand','Brand/category interest and regional changes from Google Trends. Website values are normalized 0–100, not absolute search counts. [6]'),
('SEO and paid search','Indexed content, search rankings and provider-estimated keywords/traffic/spend. Keep channel-specific estimates separate from Meta.'),
('Other advertising channels','Creative observations from other platforms’ public libraries where supported. Each library has different coverage; this is an extension to validate.'),
('Technology evidence','Visible ecommerce platform, analytics tags and forms. A detected tag does not prove working tracking or disclose events.'),
('Company scale signals','Public locations, product launches, job openings and company-announced expansion/funding. These are contextual evidence, not inferred revenue.'),
('Financial disclosures','Revenue, reporting period, entity scope and marketing expense if actually disclosed in filings or investor reports. For Malaysian listed companies, start with Bursa/issuer reports. [10]'),
('Estimated revenue','Provider model or an explicit traffic × conversion-rate × order-value scenario. Label every assumed input; this is not observed company sales.')])
h('Provider claims need verification')
p('Similarweb’s Ad Intelligence page describes modeled impressions/spend, while its FAQ still marks social coverage and API access as forthcoming. Do not treat a product headline or an old launch target as proof that Meta ad-level estimates are available now. Request coverage and sample exports for your actual businesses. [7]')
p('A traffic spike that coincides with new Meta ads is correlation. PR, email, search, marketplaces, offline sales and seasonality may explain the same change.')

page('5 Conditional disclosures and private metrics')
h('Public disclosure exceptions')
table([
('EU-delivered ads','Estimated EU reach; age/country/gender breakdown; selected target ages, gender and locations; beneficiary/payer. Preserve exact field definitions. [1]'),
('UK-delivered ads','Public archive coverage exists. Validate the particular UI/API disclosure fields instead of assuming every EU field is present. [1, 2]'),
('Political or issue ads','Spend and impression ranges, currency, sponsor/byline, audience-size range, demographic and regional distribution where supported. [1]'),
('Other regional exceptions','For example, Meta documents specific Brazilian political-ad reach/targeting fields. Enable only for eligible records. [1]')])
p('Reach is not impressions. Meta’s introductory prose and detailed field labels can differ; preserve the returned field name and its definition. In particular, eu_total_reach is defined as estimated combined EU reach. Do not relabel it as impressions or global reach. Ranges need lower/upper bounds and an open-ended flag.')
h('With authorization to the advertiser account')
p('Meta’s Insights SDK exposes the following families. Availability varies by API version, objective, event tracking, attribution and breakdown compatibility; a listed field is not guaranteed to populate. These are for your own or consenting clients’ accounts. [8]')
table([
('Delivery and cost','Spend, impressions, reach, frequency, CPM; all/link/outbound/unique clicks where supported, CTR and CPC.'),
('Site and conversion events','Landing-page views, content views, add-to-cart, checkout, purchases, leads and cost per selected action.'),
('Value','Attributed action values and purchase ROAS where supported. Attribution windows and modeled results must be recorded.'),
('Video','Plays, supported timed views, ThruPlays, average watch time and 25/50/75/95/100 percent viewing metrics.'),
('Engagement','Reported post engagement and reactions/comments/shares; messaging and other action types where supported.'),
('Breakdowns','Daily period, campaign/ad set/ad, placement/platform, device, country/region, age/gender where supported. Some combinations are restricted.'),
('Account configuration','Permitted campaign objective, ad-set targeting, optimization, budget, bid strategy, delivery status and creative IDs; these use account objects as well as Insights.')])
p('Orders, net revenue after refunds, margin, qualified leads, pipeline, customer acquisition cost, repeat purchase and lifetime value require your commerce/CRM/cost records. Ad-attributed revenue is not audited revenue or incremental revenue.')

page('6 Formulas and interpretation')
p('These are proposed dashboard calculations. All numerator/denominator inputs must share the same period, geography, currency and definition. If required inputs are missing, show Unknown.')
table([
('CTR and CPC','CTR = chosen clicks / impressions × 100%. CPC = spend / the same chosen clicks. Label all-click, link-click or outbound-click.'),
('CPM and frequency','CPM = spend / impressions × 1,000. Frequency = impressions / reach. Do not average ratios without their weights.'),
('CPA and ROAS','CPA = spend / attributed acquisitions. ROAS = attributed purchase value / spend; not profit.'),
('Video analysis','Example hook ratio = chosen early views / impressions. Hold ratio = chosen later views / early views. Store the precise definitions; these are not universal standards.'),
('Observed ad-count change','Current observed active IDs minus the prior comparable snapshot. Check collection completeness before attributing the change to the business.'),
('Offer and format mix','Number of ads with the selected attribute / eligible classified ads. Show unclassified counts and overlapping categories.'),
('Spend scenario','If sourced reach exists: reach × assumed frequency × assumed CPM / 1,000. If impressions exist: impressions × assumed CPM / 1,000.'),
('Revenue scenario','Visits × assumed conversion rate × assumed average order value. This excludes unmodeled channels; do not attribute the result to Meta.'),
('Connected business efficiency','Define total revenue / total marketing spend, or revenue / Meta spend, explicitly. Different denominators are different metrics.')])
h('Spend example without false precision')
p('Illustration only: 10,000 reached accounts, assumed frequency 1–3 and CPM RM10–30 yield RM100–900 for that same reach scope. This is a sensitivity range, not a confidence interval. If reach is unavailable, ad age alone supplies no defensible spend range.')
h('Time series rules')
p('Keep cumulative snapshots separate from incremental daily results. Subtract successive cumulative values only when scope is unchanged and corrections are handled; a negative change may be a revision. Do not add overlapping demographic slices, daily unique reach or repeated cumulative records.')
p('Compare creative cohorts with similar market, launch period and observation coverage. Active ads have unfinished lifetimes; missing records do not prove termination. A longevity score should rank research attention, never claim profitability.')

page('7 Collection and dashboard plan')
table([
('Official Library API','Use for verified eligible coverage and documented fields. It is not a universal feed of MY/SG commercial ads. Store API version and pagination results.'),
('Collector provider','Validate actual Pages, country filters, variants, snapshot history, deletion policy, media retention and costs against a sample. Preserve the provider’s raw field names.'),
('Browser or Playwright','Useful for visible-page inspection and validating collector output. It cannot reveal private CTR, spend or sales. Layout changes and access restrictions affect reliability.'),
('Own-account API','Add Meta authorization for actual account reporting; connect commerce/CRM separately to reconcile business outcomes.'),
('Refresh proposal','Daily public-ad snapshots, weekly website/offer checks and monthly contextual/financial updates are a starting design, not a schedule already running.')],('Collection method','Role'))
h('Evidence stored with every metric')
p('Entity ID; source URL/provider; raw field name/value; evidence class; capture timestamp; source update time; period start/end; geography; unit/currency; lower/upper bounds; cumulative versus incremental; definition; API/schema version; quality flag; and missing reason. Store counts separately from estimates and your own annotations.')
h('Proposed database entities')
p('Businesses → advertiser Pages → ads → creative variants/assets → observations. Add metric observations, landing-page snapshots, business metrics, research annotations and collection runs. Keep business-to-Page mapping evidence and review state. These additional tables are proposed; they are not all in the current dashboard.')
h('Recommended implementation order')
p('First: verified Page mapping, collection coverage, creative snapshots, new-ad alerts, deduplication, offer history and longevity observations. Second: transcription/OCR, creative families, copy/angle analysis and landing-page changes. Third: regional disclosure metrics and transparent scenarios. Fourth: authorized account and commerce reporting.')
h('What the current dashboard already stores')
p('Company identity/market/notes; saved ad IDs/copy/format/status/start date/URLs; pins and collections; hook/angle/notes/hypothesis; supplied reach/spend/source and scenario inputs; manual/import observations; activity; JSON export. Live collection, automatic AI analysis, business financial enrichment and account reporting are not connected.')
h('Where the information lives')
p('This catalog and the research guide are local documents in Documents/MetaAdsResearch. The published dashboard stores user-entered records in its hosted database; Data & connections → Export workspace downloads those records. There is no automatically collected competitor dataset yet.')

page('8 Sources and verification notes')
p('Checked 3 October 2026. Source availability varies; the list below supports the data boundaries, not a guarantee that every field will appear for every company. Proposed analytics and database design are our recommendations.')
sources=[
('1 Meta Ad Library API detailed fields and scope','https://br-fr.facebook.com/ads/library/api/?source=onboarding'),
('2 Meta Ad Library public archive coverage and update timing','https://business.facebook.com/ads/library/?active_status=all&ad_type=political_and_issue_ads&country=KR'),
('3 Apify Facebook Ads Library Scraper output examples','https://apify.com/apify/facebook-ads-scraper'),
('4 Meta Page transparency explanation','https://about.fb.com/news/2019/03/a-better-way-to-learn-about-ads/'),
('5 Similarweb marketing channel metrics','https://support.similarweb.com/hc/en-us/articles/4403644924561-Segments-Marketing-Channels'),
('6 Google Trends normalization and limitations','https://support.google.com/trends/answer/4365533?hl=en'),
('7 Similarweb Ad Intelligence scope and FAQ','https://support.similarweb.com/hc/en-us/articles/24435120044957-Ad-Intelligence-Overview'),
('8 Meta maintained Insights SDK field definitions','https://github.com/facebook/facebook-python-business-sdk/blob/main/facebook_business/adobjects/adsinsights.py'),
('9 Meta explanation of ad auction inputs','https://about.fb.com/wp-content/uploads/2023/01/Toward_fairness_in_personalized_ads.pdf'),
('10 Bursa investor relations example of financial disclosures','https://bursa.listedcompany.com/')]
for label,url in sources:link(label,url)
h('Pilot acceptance checks')
p('For a small set of businesses, compare collected IDs and creative variants against the Library under the same filters; record missing fields and scrape failures. Verify one EU/UK example before enabling regional panels. Confirm that provider spend labels disclose a methodology. Use an authorized sample account to test Insights fields and attribution settings before offering private performance charts.')
p('This catalog does not revise the live dashboard, purchase a provider subscription, create a tracking schedule or claim access to competitors’ private accounts.')
path=out/'Meta_Ads_Metrics_and_Retrievable_Data_Catalog.docx';d.save(path);print(path)
